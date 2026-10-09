"""Readable Excel and safe CSV snapshots. No external formulas or macros."""
from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from .sources import SOURCES

DETAILS = [
    ('url', 'URL форума'), ('title', 'Название страницы'), ('group', 'Группа'),
    ('score', 'Баллы 0–100'), ('coverage', 'Покрытие, %'),
    ('crux_lookup', 'CrUX: проверенный origin'), ('crux_rank', 'CrUX: верхняя граница диапазона'),
    ('crux_status', 'CrUX: наличие данных'), ('majestic_lookup', 'Majestic: проверенный домен'),
    ('majestic_rank', 'Majestic: место'), ('majestic_status', 'Majestic: наличие данных'),
    ('ref_subnets', 'Ссылающиеся подсети'), ('ref_ips', 'Ссылающиеся IP'),
    ('posts', 'Сообщений за всё время'), ('topics', 'Тем за всё время'), ('members', 'Регистраций за всё время'),
    ('posts_per_day', 'Чистый прирост сообщений / день'), ('growth_days', 'Интервал замеров, дней'),
    ('growth_since', 'Предыдущий замер, UTC'), ('last_activity', 'Последняя замеченная дата сообщения, UTC'),
    ('engine', 'Движок'), ('status', 'Доступ'), ('rank_scope', 'Область внешнего рейтинга'),
    ('final_url', 'Конечный URL'), ('checked_at', 'Проверка, UTC'), ('components', 'Расчёт балла'), ('notes', 'Примечания'),
]


def clean(value):
    if not isinstance(value, str):
        return value
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value).encode('utf-8', errors='replace').decode('utf-8')[:12000]


def csv_safe(value):
    value = clean(value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def put(ws, row: int, values):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, clean(value))
        if isinstance(value, str):
            cell.data_type = 's'  # Titles beginning with '=' must remain plain text.
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        cell.font = Font(name='Calibri', size=11)


def table(ws, name: str, start: int, end: int, columns: int):
    tab = Table(displayName=name, ref=f'A{start}:{get_column_letter(columns)}{end}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)
    ws.row_dimensions[start].height = 32
    for cell in ws[start][:columns]:
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='244761')
    ws.sheet_view.showGridLines = False


def export_report(result: dict, output: Path) -> dict[str, Path]:
    output.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S_%f')
    prefix = 'demo' if result.get('demo') else 'forums'
    xlsx = output / f'{prefix}_{stamp}.xlsx'
    csv_path = output / f'{prefix}_{stamp}.csv'
    rows = result['rows']
    wb = Workbook()
    main = wb.active
    main.title = 'Главная'
    metrics = wb.create_sheet('Метрики')
    sources = wb.create_sheet('Источники')
    guide = wb.create_sheet('Как читать')
    title = 'ДЕМОНСТРАЦИЯ — ВЫМЫШЛЕННЫЕ ДАННЫЕ' if result.get('demo') else 'Приоритет проверки форумов'
    put(main, 1, [title]); main.merge_cells('A1:I1')
    main['A1'].font = Font(name='Calibri', size=18, bold=True, color='244761')
    main.row_dimensions[1].height = 32
    unknown = sum(r.get('coverage', 0) < 60 for r in rows)
    put(main, 2, [f'Форумов: {len(rows)}    |    Нужны дополнительные данные: {unknown}    |    Проверка: {result["checked_at"]}'])
    main.merge_cells('A2:I2')
    put(main, 3, ['Это не число посетителей. Баллы сравнивают только этот список; A означает «проверить первым», а не «покупать рекламу».'])
    main.merge_cells('A3:I3'); main.row_dimensions[3].height = 30
    warning = ('ПРОВЕРКА ПРЕРВАНА: необработанные строки сохранены. ' if result.get('cancelled') else '')
    warning += ' | '.join(result.get('warnings', []))
    errors = [f'{kind}: {meta.get("warning", "ошибка")}' for kind, meta in result.get('sources', {}).items() if meta.get('status') != 'ok']
    warning += (' | ' if warning and errors else '') + ' | '.join(errors)
    put(main, 4, [warning or 'Пустая ячейка = нет данных. Полные показатели, источники и методика — на соседних листах.'])
    main.merge_cells('A4:I5'); main.row_dimensions[4].height = 28
    headers = ['Группа', 'Форум / ссылка', 'Баллы', 'Покрытие, %', 'Доступ', 'Что проверить',
               'Аудитория: 1–5 (вручную)', 'Цена (вручную)', 'Решение (вручную)']
    put(main, 7, headers)
    for number, row in enumerate(rows, start=8):
        put(main, number, [row.get('group'), row.get('label') or row.get('url'), row.get('score'), row.get('coverage'),
                           row.get('status'), row.get('notes'), None, None, None])
        if row.get('valid') and row.get('url', '').startswith(('https://', 'http://')):
            main.cell(number, 2).hyperlink = row['url']
            main.cell(number, 2).font = Font(name='Calibri', size=11, color='0563C1', underline='single')
        group = row.get('group', '')
        color = ('E2F0D9' if group.startswith('A') else 'FFF2CC' if group.startswith('B') else
                 'FCE4D6' if group.startswith('C') else 'E7E6E6')
        main.cell(number, 1).fill = PatternFill('solid', fgColor=color)
        main.cell(number, 3).number_format = '0.0'
        main.row_dimensions[number].height = 75
    if rows:
        table(main, 'ForumPriority', 7, 7 + len(rows), 9)
    for i, width in enumerate([25, 38, 10, 13, 29, 60, 18, 16, 22], start=1):
        main.column_dimensions[get_column_letter(i)].width = width
    main.freeze_panes = 'C8'
    fit = DataValidation(type='whole', operator='between', formula1=1, formula2=5, allow_blank=True)
    fit.errorTitle = 'Оценка аудитории'; fit.error = 'Введите число от 1 до 5.'; fit.showErrorMessage = True
    main.add_data_validation(fit); fit.add(f'G8:G{max(8, 7 + len(rows))}')
    decision = DataValidation(type='list', formula1='"Запросить статистику,Тестовое размещение,Отказ,Обсудить"', allow_blank=True)
    main.add_data_validation(decision); decision.add(f'I8:I{max(8, 7 + len(rows))}')
    put(metrics, 1, [title for _, title in DETAILS])
    for number, row in enumerate(rows, start=2):
        put(metrics, number, [row.get(key) for key, _ in DETAILS])
        metrics.row_dimensions[number].height = 60
    if rows:
        table(metrics, 'ForumMetrics', 1, len(rows) + 1, len(DETAILS))
    metrics.freeze_panes = 'C2'
    for i, (key, _) in enumerate(DETAILS, start=1):
        metrics.column_dimensions[get_column_letter(i)].width = 40 if key in ('url', 'notes', 'components', 'final_url', 'title') else 24
    put(sources, 1, ['Источник', 'Состояние', 'Данные / дата снимка', 'Загружено, UTC', 'URL выгрузки', 'Лицензия / происхождение'])
    for index, (kind, settings) in enumerate(SOURCES.items(), start=2):
        meta = result.get('sources', {}).get(kind, {})
        status = 'Демонстрационные числа, не загружались' if result.get('demo') else meta.get('status', 'Не использовался')
        if meta.get('stale'):
            status += ' — старый кэш'
        put(sources, index, [settings['name'], status,
                            meta.get('snapshot_updated') or meta.get('last_modified') or 'Дата снимка не установлена',
                            meta.get('downloaded_at'), meta.get('source_url') or settings['url'], settings['license']])
        sources.row_dimensions[index].height = 72
    put(sources, 4, ['Публичные страницы', 'Только разрешённое robots.txt', 'Даты проверки — на листе «Метрики»',
                     result['checked_at'], 'URL каждой страницы — на листе «Метрики»', 'Счётчики опубликованы самими форумами; не проверены администрацией инструмента'])
    put(sources, 6, ['CrUX: дата обновления файла снимка не обязательно совпадает с месяцем, за который собрана статистика.'])
    sources.merge_cells('A6:F7')
    for i, width in enumerate([36, 38, 35, 30, 60, 55], start=1):
        sources.column_dimensions[get_column_letter(i)].width = width
    table(sources, 'DataSources', 1, 4, 6); sources.freeze_panes = 'A2'
    instructions = [
        ('Назначение', 'Отчёт помогает выбрать, какие форумы проверять первыми. Посещения, уникальные посетители и ожидаемые продажи НЕ рассчитываются.'),
        ('Баллы', 'Процентили среди строк с доступным показателем: меньший CrUX/место Majestic лучше, больший чистый прирост сообщений лучше. Одинаковые значения получают одинаковые баллы. Единственное известное значение получает нейтральные 50 баллов.'),
        ('Веса', 'CrUX 50%; Majestic 30%; прирост сообщений 20%. Балл — средневзвешенное только доступных компонентов. Состав формулы указан для каждой строки на листе «Метрики».'),
        ('Покрытие', 'Сумма весов доступных компонентов, НЕ вероятность правильности оценки. Нельзя считать отсутствие источника нулевой популярностью. При покрытии <60% группа «Нужны данные», даже при высоком балле.'),
        ('Группы', 'При покрытии от 60%: A — балл от 70; B — от 35 до 70; C — ниже 35. Пороги эвристические, не откалиброваны по рекламным конверсиям. Неизвестные форумы вынесены отдельно, не в C.'),
        ('CrUX', 'Диапазон популярности origin по данным части пользователей Chrome. Число 10000 означает верхнюю границу диапазона, а не точное 10000-е место. HTTP/HTTPS, www и поддомены — разные origin. Отсутствие в Top 1M ничего не доказывает о закрытом форуме.'),
        ('Majestic', 'Ссылочная известность домена по ссылающимся подсетям. RefIPs и RefSubNets не являются числом бэклинков, ссылающихся доменов или посетителей. Поддомен не получает автоматически показатель родительского домена; www сопоставляется с тем же доменом без www.'),
        ('Закрытые форумы', 'Авторизация, CAPTCHA и запреты robots.txt не обходятся. Счётчики могут быть скрыты или накручены. Проверьте статистику у администрации: активные участники за 30 дней, просмотры рекламного раздела, география и цена.'),
        ('Ссылки на раздел', 'Для /forum или ссылки с параметрами внешняя статистика всего сайта показана для справки, но не засчитывается форуму. Нельзя приписывать разделу аудиторию всего портала. Редирект на другой домен также требует проверки.'),
        ('Прирост', 'Чистое изменение счётчика сообщений между замерами за 1–90 дней, предпочтительно около 7 дней. Это не число уникальных авторов и не валовые новые сообщения: удаления/миграции влияют на разницу. Округлённые K/M-счётчики не используются.'),
        ('Ручные колонки', 'Аудитория, цена и решение на «Главной» заполняются вами. Они не меняют автоматический рейтинг. Сохраняйте заполненный файл: новый запуск создаёт новый отчёт, ручные пометки туда автоматически не переносятся.'),
        ('Повторный запуск', 'Запускайте с той же папкой результатов. history.json хранит агрегатные замеры локально. Кэш источников обновляется не чаще раза в 24 часа. Отчёт — числовой снимок расчёта, не редактируемая финансовая модель.'),
        ('Ограничения', 'Работает с публичным HTML без JavaScript. Не каждый движок/язык распознаётся. Неизвестный показатель остаётся пустым. Это локальная утилита, не защищённый сервис для размещения в Интернете.'),
    ]
    put(guide, 1, ['Раздел', 'Объяснение'])
    for i, values in enumerate(instructions, start=2):
        put(guide, i, values); guide.row_dimensions[i].height = 72
    guide.column_dimensions['A'].width = 25; guide.column_dimensions['B'].width = 115
    table(guide, 'ReportGuide', 1, len(instructions) + 1, 2); guide.freeze_panes = 'A2'
    xlsx_temp = xlsx.with_suffix('.tmp.xlsx')
    csv_temp = csv_path.with_suffix('.tmp.csv')
    try:
        wb.save(xlsx_temp)
        with csv_temp.open('w', encoding='utf-8-sig', newline='') as out:
            writer = csv.writer(out, delimiter=';')
            writer.writerow([title for _, title in DETAILS])
            for row in rows:
                writer.writerow([csv_safe(row.get(key)) for key, _ in DETAILS])
        xlsx_temp.replace(xlsx); csv_temp.replace(csv_path)
    finally:
        wb.close()
        xlsx_temp.unlink(missing_ok=True); csv_temp.unlink(missing_ok=True)
    return {'xlsx': xlsx, 'csv': csv_path}
