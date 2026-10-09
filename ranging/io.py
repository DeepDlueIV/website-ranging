"""Small TXT/CSV/XLSX input reader. Invalid rows are retained in the report."""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from .core import normalize_url

MAX_ROWS = 1000
URL_HEADERS = {'url', 'link', 'website', 'site', 'ссылка', 'ссылки', 'форум', 'сайт', 'адрес'}
NAME_HEADERS = {'name', 'label', 'название', 'название форума'}


def read_inputs(source: str | Path) -> tuple[list[dict], list[str]]:
    warnings = []
    records = []
    if isinstance(source, Path):
        if source.stat().st_size > 10_000_000:
            raise ValueError('Входной файл слишком большой: максимум 10 МБ')
        suffix = source.suffix.lower()
        if suffix == '.xlsx':
            with zipfile.ZipFile(source) as archive:
                if sum(item.file_size for item in archive.infolist()) > 60_000_000:
                    raise ValueError('Excel содержит слишком много распакованных данных')
            wb = load_workbook(source, read_only=False, data_only=False, keep_links=False)
            try:
                ws = wb['Ссылки'] if 'Ссылки' in wb.sheetnames else wb.worksheets[0]
                if ws.max_row > MAX_ROWS + 1 or ws.max_column > 100:
                    raise ValueError('Не более 1000 строк и 100 колонок. Удалите лишнее форматирование из файла.')
                for row in ws.iter_rows():
                    values = []
                    for cell in row:
                        if cell.data_type == 'f':
                            values.append('[Формула вместо ссылки — вставьте URL текстом]')
                        elif cell.hyperlink and cell.hyperlink.target:
                            values.append(cell.hyperlink.target)
                        else:
                            values.append(cell.value)
                    records.append(values)
            finally:
                wb.close()
        elif suffix in ('.txt', '.csv', '.tsv'):
            raw = source.read_bytes()
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = raw.decode('cp1251')
                warnings.append('Файл прочитан в кодировке Windows-1251.')
            if suffix == '.txt':
                records = [[line] for line in text.splitlines()]
            else:
                try:
                    dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
                except csv.Error:
                    dialect = csv.excel
                records = list(csv.reader(io.StringIO(text), dialect))
        else:
            raise ValueError('Поддерживаются TXT, CSV и XLSX. Старый XLS сохраните как XLSX.')
    else:
        records = [[line] for line in source.splitlines()]
    records = [r for r in records if r and any(v is not None and str(v).strip() for v in r)
               and not str(r[0] or '').lstrip().startswith('#')]
    if not records:
        raise ValueError('Список пуст. Вставьте ссылки по одной на строке.')
    first = [str(v or '').strip().lower() for v in records[0]]
    url_col = next((i for i, value in enumerate(first) if value in URL_HEADERS), 0)
    name_col = next((i for i, value in enumerate(first) if value in NAME_HEADERS), None)
    if any(value in URL_HEADERS for value in first):
        records = records[1:]
    if len(records) > MAX_ROWS:
        raise ValueError('За один запуск поддерживается не более 1000 ссылок.')
    rows = []
    seen = set()
    for number, record in enumerate(records, start=1):
        raw = str(record[url_col] or '').strip() if url_col < len(record) else ''
        if not raw:
            warnings.append(f'Строка {number}: пустая ссылка пропущена.')
            continue
        label = str(record[name_col] or '') if name_col is not None and name_col < len(record) else ''
        row = {'input_url': raw, 'url': raw, 'label': label, 'valid': False, 'notes': '', 'input_row': number}
        try:
            row['url'] = normalize_url(raw)
            row['valid'] = True
        except ValueError as exc:
            row['notes'] = str(exc)
        if row['valid'] and row['url'] in seen:
            warnings.append(f'Дубликат пропущен: {row["url"]}')
            continue
        if row['valid']:
            seen.add(row['url'])
        rows.append(row)
    if not rows:
        raise ValueError('Не найдено ни одной непустой строки со ссылкой.')
    return rows, warnings
