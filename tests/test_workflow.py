import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

from ranging.io import read_inputs
from ranging.engine import analyze
from ranging.report import export_report
from ranging.sources import SourceError


class WorkflowTests(unittest.TestCase):
    def test_paste_dedup_and_invalid_are_visible(self):
        rows, warnings = read_inputs('example.com\nhttps://example.com/\nfile:///bad\n# comment\n')
        self.assertEqual(len(rows), 2)
        self.assertTrue(warnings)
        self.assertFalse(rows[1]['valid'])

    def test_csv_semicolon_header(self):
        with tempfile.TemporaryDirectory() as td:
            file = Path(td)/'list.csv'
            file.write_text('Название;URL\nНаш форум;https://example.com/forum/\n', encoding='utf-8-sig')
            rows, _ = read_inputs(file)
        self.assertEqual(rows[0]['label'], 'Наш форум')
        self.assertTrue(rows[0]['url'].endswith('/forum/'))

    def test_xlsx_hyperlinks_and_formula_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            file = Path(td)/'list.xlsx'
            wb = Workbook(); ws = wb.active
            ws.append(['URL'])
            ws.append(['Open forum']); ws['A2'].hyperlink = 'https://example.com/'
            ws.append(['=HYPERLINK("https://bad.example","x")'])
            wb.save(file)
            rows, _ = read_inputs(file)
        self.assertEqual(rows[0]['url'], 'https://example.com/')
        self.assertFalse(rows[1]['valid'])

    def test_whole_pipeline_keeps_source_failures_unknown(self):
        with tempfile.TemporaryDirectory() as td:
            rows, _ = read_inputs('example.com\nother.example')
            with patch('ranging.engine.load_dataset', side_effect=SourceError('offline')):
                result = analyze(rows, Path(td), inspect_pages=False)
        self.assertEqual(len(result['rows']), 2)
        self.assertTrue(all(row['score'] is None for row in result['rows']))
        self.assertEqual(result['sources']['crux']['status'], 'error')

    def test_actual_source_ranks_flow_to_scores(self):
        def data(kind, wanted, cache, client):
            key = 'https://example.com' if kind == 'crux' else 'example.com'
            return {key: {'rank': 1000, 'ref_subnets': 50, 'ref_ips': 100}}, {'status':'ok', 'source_url':'https://source.example'}
        with tempfile.TemporaryDirectory() as td:
            rows, _ = read_inputs('example.com')
            with patch('ranging.engine.load_dataset', side_effect=data):
                result = analyze(rows, Path(td), inspect_pages=False)
        self.assertEqual(result['rows'][0]['coverage'], 80)
        self.assertEqual(result['rows'][0]['score'], 50)

    def test_export_safe_complete_and_readable(self):
        with tempfile.TemporaryDirectory() as td:
            row = {'url':'https://example.com/', 'label':'=1+1', 'title':'evil\x01',
                   'score':50, 'coverage':80, 'group':'B — второй приоритет',
                   'notes':'@SUM(A1:A2)', 'valid':True, 'status':'test'}
            result = {'rows':[row], 'sources':{}, 'warnings':[], 'checked_at':'2026-10-09T12:00:00+00:00', 'cancelled':False}
            paths = export_report(result, Path(td))
            wb = load_workbook(paths['xlsx'])
            self.assertEqual(wb.sheetnames, ['Главная', 'Метрики', 'Источники', 'Как читать'])
            for sheet in wb:
                for cells in sheet:
                    for cell in cells:
                        self.assertNotEqual(cell.data_type, 'f')
            self.assertEqual(wb['Главная'].freeze_panes, 'C8')
            self.assertTrue(wb['Главная'].tables)
            self.assertIn("'@SUM", paths['csv'].read_text(encoding='utf-8-sig'))

    def test_empty_input_clear_error(self):
        with self.assertRaises(ValueError):
            read_inputs(' # only comment\n')

if __name__ == '__main__':
    unittest.main()
